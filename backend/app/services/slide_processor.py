"""Slide processing service (Ideal Agent Stage 1): deck → per-slide PNGs + text.

Accepts .pptx and .pdf. A .pptx is rendered .pptx → PDF (LibreOffice headless) →
PNG per page (pdftoppm) — LibreOffice's direct PNG export only renders the first
slide, so the PDF hop is the reliable path. A .pdf skips straight to pdftoppm.
Text comes from python-pptx (.pptx) or pdftotext (.pdf).
"""

import logging
import shutil
import subprocess
from pathlib import Path

from pptx import Presentation

from app.config import settings
from app.services.media import MediaError, run_command
from app.services.storage import storage_service

logger = logging.getLogger(__name__)


class SlideProcessorError(Exception):
    """Raised when slide processing fails."""


class SlideProcessor:
    """Converts .pptx files to per-slide PNG images using LibreOffice headless."""

    async def convert_pptx_to_pngs(self, session_id: str, pptx_relative_path: str) -> list[Path]:
        """
        Convert a .pptx file to PNG images, one per slide (~1920px wide).

        Returns:
            Sorted list of absolute paths: slides/slide_1.png ... slide_N.png

        Raises:
            SlideProcessorError: If conversion fails.
        """
        pptx_abs = storage_service.get_absolute_path(pptx_relative_path)
        if not pptx_abs.exists():
            raise SlideProcessorError(f"PPTX file not found: {pptx_abs}")

        dirs = storage_service.ensure_session_dirs(session_id)
        slides_dir = dirs["slides"]
        for stale in slides_dir.glob("*.png"):
            stale.unlink()

        work_dir = dirs["root"] / "render"
        shutil.rmtree(work_dir, ignore_errors=True)
        work_dir.mkdir(parents=True)

        if pptx_abs.suffix.lower() == ".pdf":
            shutil.copy(pptx_abs, work_dir / "deck.pdf")
        else:
            await self._pptx_to_pdf(pptx_abs, work_dir)
        pdfs = sorted(work_dir.glob("*.pdf"))
        if not pdfs:
            raise SlideProcessorError("LibreOffice did not produce a PDF — is the .pptx valid?")
        return await self._pdf_to_pngs(pdfs[0], work_dir, slides_dir, dirs["root"])

    async def _pptx_to_pdf(self, pptx_abs: Path, work_dir: Path) -> None:
        # Separate LibreOffice profile per job so concurrent conversions don't collide
        profile = (work_dir / "lo_profile").resolve().as_uri()
        try:
            await run_command(
                [
                    settings.libreoffice_bin,
                    f"-env:UserInstallation={profile}",
                    "--headless",
                    "--convert-to", "pdf",
                    "--outdir", str(work_dir),
                    str(pptx_abs),
                ],
                timeout=180,
            )
        except MediaError as e:
            raise SlideProcessorError(f"PPTX → PDF conversion failed: {e}") from e

    async def _pdf_to_pngs(self, pdf: Path, work_dir: Path, slides_dir: Path, root: Path) -> list[Path]:
        try:
            await run_command(
                [
                    "pdftoppm", "-png",
                    "-scale-to-x", "1920", "-scale-to-y", "-1",
                    str(pdf),
                    str(work_dir / "page"),
                ],
                timeout=180,
            )
        except MediaError as e:
            raise SlideProcessorError(f"PDF → PNG conversion failed: {e}") from e

        # pdftoppm zero-pads page numbers to equal width, so a lexical sort is page order
        pages = sorted(work_dir.glob("page-*.png"))
        if not pages:
            raise SlideProcessorError("No slide images were produced from the deck")

        normalized = []
        for i, page in enumerate(pages, start=1):
            target = slides_dir / f"slide_{i}.png"
            page.rename(target)
            normalized.append(target)

        shutil.move(str(pdf), root / "slides.pdf")
        shutil.rmtree(work_dir, ignore_errors=True)
        return normalized

    def extract_slide_texts(self, pptx_relative_path: str) -> list[str]:
        """
        Extract visible text per (non-hidden) slide, titles first.

        Used as extra VLM context and as the offline script fallback. Hidden
        slides are skipped because LibreOffice leaves them out of the PDF export.
        """
        pptx_abs = storage_service.get_absolute_path(pptx_relative_path)
        if pptx_abs.suffix.lower() == ".pdf":
            return self._pdf_texts(pptx_abs)
        try:
            prs = Presentation(str(pptx_abs))
        except Exception as e:  # noqa: BLE001 — text is best-effort
            logger.warning("Could not read slide text from %s: %s", pptx_abs, e)
            return []

        texts = []
        for slide in prs.slides:
            if slide.element.get("show") == "0":
                continue
            lines: list[str] = []
            title_shape = slide.shapes.title
            if title_shape is not None and title_shape.text_frame.text.strip():
                lines.append(title_shape.text_frame.text.strip())
            title_id = title_shape.shape_id if title_shape is not None else None
            for shape in slide.shapes:
                # python-pptx builds new proxies on each access, so compare ids, not identity
                if shape.shape_id == title_id:
                    continue
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = "".join(run.text for run in para.runs).strip()
                        if text:
                            lines.append(text)
                elif shape.has_table:
                    for row in shape.table.rows:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            lines.append(" | ".join(cells))
            texts.append("\n".join(lines))
        return texts

    @staticmethod
    def _pdf_texts(pdf: Path) -> list[str]:
        """Text per page via pdftotext (pages are separated by form feeds)."""
        try:
            out = subprocess.run(
                ["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True, timeout=120, check=True
            ).stdout
        except (OSError, subprocess.SubprocessError) as e:
            logger.warning("Could not read text from %s: %s", pdf, e)
            return []
        pages = out.split("\f")
        if pages and not pages[-1].strip():
            pages = pages[:-1]
        # Collapse the column padding -layout adds, keep one line per text line
        return ["\n".join(" ".join(line.split()) for line in page.splitlines() if line.strip()) for page in pages]


# Module-level singleton
slide_processor = SlideProcessor()
